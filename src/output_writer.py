"""Çözüm çıktılarını yarışma şablonlarıyla birebir eşleşen formatta yazar."""
from __future__ import annotations

from datetime import datetime, time

import openpyxl
import pandas as pd

TALEP_TAHMIN_KOLONLARI = [
    "Talep ID",
    "Tarih",
    "Talep Tamamlama Saati",
    "Çıkış Transfer Merkezi",
    "Varış Transfer Merkezi",
    "Tahmin Edilen Desi",
]

TASIMA_PLANI_KOLONLARI = [
    "Araç ID",
    "Araç Tipi",
    "Araç türü",
    "Çıkış Transfer Merkezi",
    "Varış Transfer Merkezi",
    "Çıkış Tarihi",
    "Çıkış Saati",
    "Varış Tarihi",
    "Varış Saati",
    "Talep ID",
    "Taşınan Desi",
    "Yolculuk süresi",
    "Varış elleçleme süresi",
    "Çıkış Elleçleme süresi",
    "SLA cezası",
    "Toplam maliyet",
]


def _parse_saat(saat_str: str) -> time:
    hh, mm = saat_str.split(":")
    return time(int(hh), int(mm))


def write_talep_tahmini(forecast: pd.DataFrame, path) -> None:
    """forecast kolonları: talep_id, tarih, talep_tamamlanma_saati, cikis, varis, tahmin_desi"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(TALEP_TAHMIN_KOLONLARI)

    for row in forecast.itertuples(index=False):
        tarih_str = pd.Timestamp(row.tarih).strftime("%d.%m.%Y")
        ws.append(
            [
                row.talep_id,
                tarih_str,
                _parse_saat(row.talep_tamamlanma_saati),
                row.cikis,
                row.varis,
                row.tahmin_desi,
            ]
        )
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = "h:mm"

    wb.save(path)


def write_tasima_plani(plan: pd.DataFrame, path) -> None:
    """plan kolonları TASIMA_PLANI_KOLONLARI ile birebir aynı isimde olmalı (Türkçe)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(TASIMA_PLANI_KOLONLARI)

    for row in plan.itertuples(index=False):
        ws.append(list(row))

    cikis_saat_col = TASIMA_PLANI_KOLONLARI.index("Çıkış Saati") + 1
    varis_saat_col = TASIMA_PLANI_KOLONLARI.index("Varış Saati") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=cikis_saat_col).number_format = "hh:mm"
        ws.cell(row=r, column=varis_saat_col).number_format = "hh:mm"

    wb.save(path)
